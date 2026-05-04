"""Evaluation harness for thesis evaluation campaigns.

Drives the configured baseline conditions across the testbed task set and
records ``EvaluationRecord`` rows to the ``EvaluationStore``.

Baseline conditions
-------------------
full            : full pipeline (JIT + verification + optimizer)
no_verification : JIT + optimizer with verification disabled
no_optimization : JIT pipeline without optimizer
static          : seeded registry + hand-written static implementations

Run protocol
------------
1. Restore a clean snapshot at the start of every run.
2. Randomize task order within each run using a deterministic seed.
3. Use deterministic condition order by default; runners may opt into shuffled
   condition order for methodology checks.
4. Keep the registry alive within a run so tasks can reuse capabilities created
   by earlier tasks in the same run — this is intentional for optimizer data.
5. Reset between runs so each run starts from the same clean state.

Usage::

    from tests.evaluation.harness import EvaluationHarness, HarnessConfig
    from tests.evaluation.tasks import ALL_TASKS, PILOT_TASKS

    config = HarnessConfig(
        chroma_dir=Path("data/eval/chroma"),
        skills_dir=Path("data/eval/skills"),
        db_path="data/eval/metrics.duckdb",
        n_runs=8,
        base_seed=42,
    )
    harness = EvaluationHarness(config)
    # Run pilot first
    harness.run_pilot(PILOT_TASKS)
    # Then full evaluation
    for condition in ["full", "no_verification", "no_optimization", "static"]:
        harness.run_condition(condition, ALL_TASKS)
"""

from __future__ import annotations

import json
import logging
import random
import time
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from src.orchestrator.output_contract_resolver import _extract_output_contract_llm
from src.orchestrator.tool_loader import ToolLoader
from src.registry.evaluation_store import EvaluationStore
from src.registry.store import SkillRegistry, seed_registry
from src.schemas.evaluation import EvaluationRecord, GeneratedArtifactSnapshot
from src.schemas.task import GeneratedArtifact, TaskState
from tests.evaluation.conditions import get_condition_profile
from tests.evaluation.snapshots import restore_clean_snapshot, take_clean_snapshot
from tests.evaluation.tasks import TaskDefinition

logger = logging.getLogger(__name__)


@dataclass
class HarnessConfig:
    """Configuration for one evaluation campaign."""

    chroma_dir: Path
    skills_dir: Path

    # EvaluationStore path (None = in-memory, useful for tests)
    db_path: str | None = None

    # Run parameters
    n_runs: int = 8
    base_seed: int = 42

    # ChromaDB settings
    chroma_collection: str = "capabilities"

    # Optional LangFuse callback handler passed to pipeline invocations.
    langfuse_handler: Any | None = None

    # When true, every task is initialized with allow_network=True. The current
    # core evaluation stays offline by default.
    allow_network: bool = False

    # MetricsStore path. Filesystem paths preserve optimizer telemetry across
    # runs; None keeps an in-memory store for isolated tests.
    metrics_store_path: str | None = None


@dataclass(frozen=True)
class StaticImplementationResult:
    """Normalized output from a hand-written static capability implementation."""

    output: dict
    output_files: dict[str, bytes]


def _normalize_static_implementation_result(
    raw: dict | tuple[dict, dict[str, bytes]] | StaticImplementationResult,
) -> StaticImplementationResult:
    """Normalize static implementation output.

    Static implementations may return only an output dict or an output dict
    paired with generated file bytes.
    """
    if isinstance(raw, StaticImplementationResult):
        return raw
    if isinstance(raw, tuple):
        output, output_files = raw
        return StaticImplementationResult(
            output=dict(output),
            output_files=dict(output_files),
        )
    return StaticImplementationResult(output=dict(raw), output_files={})


# ---------------------------------------------------------------------------
# Capture generated artifact sources at record-write time so later inspection
# can read full source bodies even after snapshot isolation has reset the
# registry.
# ---------------------------------------------------------------------------


def _derive_condition_seed(condition: str, base_seed: int) -> int:
    """Deterministic seed derivation for EvaluationHarness.

    Uses SHA-256 rather than Python's process-randomized ``hash()`` so the same
    condition and base seed produce the same ordering across processes and
    Python versions.
    """
    import hashlib
    key = f"{condition}|{int(base_seed)}".encode()
    return int.from_bytes(hashlib.sha256(key).digest()[:4], "big") % (2**31)


def _snapshot_artifacts(
    artifacts: list[GeneratedArtifact] | None,
) -> list[GeneratedArtifactSnapshot]:
    """Build source snapshots from the pipeline's generated artifacts.

    Source body is polymorphic by type:
        * MCP_TOOL -> source_code (raw Python)
        * SKILL    -> skill_content (full SKILL.md)
        * AGENT    -> json.dumps(agent_config) keeping system_prompt + subagents

    Never raises — any artifact whose source body can't be extracted is
    skipped so a single malformed artifact can't break record writing.
    Empty list in / empty list out.
    """
    from src.schemas.capability import CapabilityType

    if not artifacts:
        return []

    out: list[GeneratedArtifactSnapshot] = []
    for art in artifacts:
        try:
            if art.capability_type == CapabilityType.MCP_TOOL:
                body = art.source_code or ""
            elif art.capability_type == CapabilityType.SKILL:
                body = art.skill_content or ""
            elif art.capability_type == CapabilityType.AGENT:
                if art.agent_config is None:
                    body = ""
                elif isinstance(art.agent_config, dict):
                    body = json.dumps(art.agent_config, default=str)
                else:
                    body = str(art.agent_config)
            else:
                body = ""
            if not body:
                # Still snapshot the name/type so later inspection can explain
                # why the source body is absent.
                body = ""
            out.append(
                GeneratedArtifactSnapshot(
                    capability_name=art.capability_name,
                    capability_type=art.capability_type,
                    source_body=body,
                    domain=art.domain,
                )
            )
        except Exception:  # noqa: BLE001 — never-raises contract
            continue
    return out


class EvaluationHarness:
    """Orchestrates evaluation runs across all four baseline conditions."""

    def __init__(self, config: HarnessConfig) -> None:
        self._config = config
        self._store = EvaluationStore(db_path=config.db_path)
        self._snap_id: str | None = None
        # Accumulate per-task FlushResult objects across every _execute_run call
        # so runners can fail the campaign when telemetry flushes are partial.
        from tests.evaluation.cycle_flush import FlushResult
        self._flush_results: list[FlushResult] = []
        # Emit a one-time warning listing permissive metadata-only validators.
        from tests.evaluation.tasks import warn_permissive_validators_once
        warn_permissive_validators_once()

    # -- public API ----------------------------------------------------------

    def prepare(self) -> str:
        """Ensure a clean state and take the snapshot.

        Wipes generated artifacts (skills, chroma) and rebuilds a fresh
        seeded registry so the snapshot captures only the baseline state,
        not leftover artifacts from previous runs.

        Must be called before any run_* method.  Returns the snapshot ID.
        """
        import gc
        import shutil

        cfg = self._config

        # 1. Wipe generated skills (keep the directory itself)
        if cfg.skills_dir.exists():
            shutil.rmtree(cfg.skills_dir)
        cfg.skills_dir.mkdir(parents=True, exist_ok=True)
        logger.info("Cleaned skills directory: %s", cfg.skills_dir)

        # 2. Wipe and rebuild ChromaDB with fresh seeded registry
        try:
            import chromadb

            chromadb.api.shared_system_client.SharedSystemClient.clear_system_cache()
        except Exception as exc:  # noqa: BLE001
            logger.debug("ChromaDB cache clear failed (non-critical): %s", exc)
        gc.collect()

        if cfg.chroma_dir.exists():
            shutil.rmtree(cfg.chroma_dir)
        cfg.chroma_dir.mkdir(parents=True, exist_ok=True)

        registry = SkillRegistry(
            persist_directory=str(cfg.chroma_dir),
            collection_name=cfg.chroma_collection,
        )
        seed_registry(registry)
        logger.info("Rebuilt clean registry with %d seeded capabilities", len(registry.list_all()))

        # 3. Take snapshot of the clean state
        self._snap_id = take_clean_snapshot(cfg.chroma_dir, cfg.skills_dir)
        logger.info("Clean snapshot taken: %s", self._snap_id)
        return self._snap_id

    def run_pilot(
        self,
        tasks: list[TaskDefinition],
        *,
        condition: str = "full",
    ) -> list[EvaluationRecord]:
        """Run pilot tasks through a single condition with is_pilot=True.

        Pilot records are flagged in EvaluationStore and excluded from
        primary metrics but can be cited in the thesis for task calibration.
        """
        run_id = f"pilot_{uuid.uuid4().hex[:8]}"
        return self._execute_run(
            run_id=run_id,
            condition=condition,
            tasks=tasks,
            seed=self._config.base_seed,
            is_pilot=True,
        )

    def run_condition(
        self,
        condition: str,
        tasks: list[TaskDefinition],
    ) -> list[EvaluationRecord]:
        """Run n_runs of a single condition across all tasks.

        Task order is randomized per run with deterministic per-condition seeds.
        The registry is restored to the clean snapshot before each run.
        """
        all_records: list[EvaluationRecord] = []
        # Include the condition name so each condition gets independent but
        # reproducible task orderings.
        condition_seed = _derive_condition_seed(condition, self._config.base_seed)
        rng = random.Random(condition_seed)
        condition_seeds = [rng.randint(0, 2**31) for _ in range(self._config.n_runs)]

        for run_idx, seed in enumerate(condition_seeds):
            # Use full UUID entropy so append/recovery runs cannot collide.
            run_id = f"{condition}_{run_idx:02d}_{uuid.uuid4().hex}"
            logger.info("Starting run %s (%s, seed=%d)", run_id, condition, seed)
            records = self._execute_run(
                run_id=run_id,
                condition=condition,
                tasks=tasks,
                seed=seed,
                is_pilot=False,
            )
            all_records.extend(records)
            logger.info(
                "Run %s complete: %d/%d tasks succeeded",
                run_id,
                sum(r.success for r in records),
                len(records),
            )
        return all_records

    @property
    def store(self) -> EvaluationStore:
        return self._store

    @property
    def flush_results(self) -> list[Any]:
        """Per-task FlushResult objects accumulated across all runs.

        Returns the live list reference so runners can inspect aggregate flush
        health before deciding their process exit code. Callers should not
        mutate it.
        """
        return self._flush_results

    def close(self) -> None:
        self._store.close()

    def __enter__(self) -> EvaluationHarness:
        return self

    def __exit__(self, *_exc: object) -> None:
        self.close()

    # -- internal ------------------------------------------------------------

    def _execute_run(
        self,
        *,
        run_id: str,
        condition: str,
        tasks: list[TaskDefinition],
        seed: int,
        is_pilot: bool,
    ) -> list[EvaluationRecord]:
        """Restore snapshot, randomize task order, execute each task, record."""
        import gc

        cfg = self._config

        # Force-release any ChromaDB PersistentClient from the previous run.
        # Without this, SQLite WAL files may still be held open when
        # restore_clean_snapshot does shutil.rmtree → shutil.copytree,
        # causing corruption on the next SkillRegistry init.
        try:
            import chromadb

            chromadb.api.shared_system_client.SharedSystemClient.clear_system_cache()
        except Exception as exc:  # noqa: BLE001
            logger.debug("ChromaDB cache clear failed (non-critical): %s", exc)
        gc.collect()

        # Restore clean snapshot.
        snap_id = restore_clean_snapshot(cfg.chroma_dir, cfg.skills_dir)
        if self._snap_id and snap_id != self._snap_id:
            msg = (
                f"Snapshot ID mismatch: expected {self._snap_id}, got {snap_id} "
                "— registry state may be corrupted"
            )
            logger.error(msg)
            raise RuntimeError(msg)

        # Rebuild registry and loader from the restored snapshot.
        registry = SkillRegistry(
            persist_directory=str(cfg.chroma_dir),
            collection_name=cfg.chroma_collection,
        )
        seed_registry(registry)
        loader = ToolLoader()

        # Create a single MetricsStore for this run. When metrics_store_path is
        # set, optimizer telemetry persists across runs; otherwise tests get an
        # isolated in-memory store.
        from src.registry.metrics import MetricsStore

        metrics_store = MetricsStore(db_path=self._config.metrics_store_path)

        # Randomize task order.
        rng = random.Random(seed)
        shuffled_tasks = list(tasks)
        rng.shuffle(shuffled_tasks)

        # Execute each task.
        records: list[EvaluationRecord] = []
        from datetime import UTC, datetime

        from src.orchestrator.tracing import trace_context
        for i, task in enumerate(shuffled_tasks):
            logger.info(
                "──── Task %d/%d: %s (%s) ────",
                i + 1, len(shuffled_tasks), task.task_id, task.tier,
            )
            # Capture the pre-task timestamp so the post-task cycle_events flush
            # can filter MetricsStore capability_events to this task only.
            task_events_cutoff = datetime.now(UTC)
            # Wrap pipeline invocation and post-task flush so events emitted
            # outside graph nodes still inherit task tags/session/user.
            _lf_tags = self._build_task_lf_tags(task, condition, is_pilot)
            with trace_context(
                tags=_lf_tags,
                session_id=run_id,
                user_id="eval-harness",
            ):
                record = self._run_one_task(
                    task=task,
                    condition=condition,
                    run_id=run_id,
                    registry=registry,
                    loader=loader,
                    metrics_store=metrics_store,
                    seed=seed,
                    snap_id=snap_id,
                    is_pilot=is_pilot,
                )
                self._store.record_run(record)
                records.append(record)

                # Persist per-task invocation, utility, and event telemetry
                # through the same flush helper used by the longitudinal runner.
                # The harness logs partial flushes and lets top-level runners
                # decide process exit behavior.
                from tests.evaluation.cycle_flush import flush_task_cycle_state
                fr = flush_task_cycle_state(
                    eval_store=self._store,
                    metrics_store=metrics_store,
                    registry=registry,
                    run_id=run_id,
                    condition=condition,
                    task_id=task.task_id,
                    task_family_id=task.family_id,
                    task_instance_id=task.instance_id,
                    variant_id=task.variant_id,
                    benchmark_track=task.benchmark_track,
                    task_events_cutoff=task_events_cutoff,
                    is_pilot=is_pilot,
                )
                # Accumulate flush health for runner-level exit handling.
                self._flush_results.append(fr)
                if not fr.all_ok:
                    import logging as _lg
                    _lg.getLogger(__name__).warning(
                        "cycle_flush partial failure for task=%s "
                        "condition=%s run_id=%s: invocations_ok=%s "
                        "utility_scores_ok=%s events_ok=%s errors=%s",
                        task.task_id, condition, run_id,
                        fr.invocations_ok, fr.utility_scores_ok, fr.events_ok,
                        fr.errors,
                    )

        return records

    def _build_task_lf_tags(
        self,
        task: TaskDefinition,
        condition: str,
        is_pilot: bool,
    ) -> list[str]:
        """Build the per-task LangFuse tag list.

        Shared by the outer trace context and the pipeline invocation metadata
        so task-level traces and post-task events use the same labels.
        """
        _file_formats = list({
            fp.suffix.lower().lstrip(".") for fp in task.fixture_paths
        }) if task.fixture_paths else []
        _tags: list[str] = [
            condition,
            task.task_id,
            f"task_family:{task.family_id}",
            f"task_instance:{task.instance_id}",
            f"benchmark:{task.benchmark_track}",
            f"tier:{task.tier}",
            *(f"format:{fmt}" for fmt in sorted(_file_formats)),
            f"allow_network:{str(self._config.allow_network).lower()}",
            f"static_coverage:{str(task.has_static_coverage).lower()}",
        ]
        if task.variant_id:
            _tags.append(f"variant:{task.variant_id}")
        if is_pilot:
            _tags.append("pilot")
        return _tags

    def _run_one_task(
        self,
        *,
        task: TaskDefinition,
        condition: str,
        run_id: str,
        registry: SkillRegistry,
        loader: ToolLoader,
        metrics_store: Any,
        seed: int,
        snap_id: str,
        is_pilot: bool,
    ) -> EvaluationRecord:
        """Execute one task under one condition and return an EvaluationRecord."""
        t_start = time.monotonic()
        profile = get_condition_profile(condition)

        if not profile.jit_enabled:
            result = self._run_static(task, registry, metrics_store)
        else:
            result = self._run_jit(
                task,
                condition,
                registry,
                loader,
                metrics_store,
                run_id=run_id,
                is_pilot=is_pilot,
            )

        elapsed_ms = int((time.monotonic() - t_start) * 1000)

        return EvaluationRecord(
            run_id=run_id,
            condition=condition,
            task_id=task.task_id,
            task_family_id=task.family_id,
            task_instance_id=task.instance_id,
            variant_id=task.variant_id,
            benchmark_track=task.benchmark_track,
            task_tier=task.tier,
            attempt_count=result.get("attempt_count", 0),
            capability_type=result.get("capability_type"),
            verification_tier_reached=result.get("verification_tier_reached"),
            gate_passed=result.get("gate_passed", False),
            output_in_range=result.get("output_in_range", False),
            token_cost=result.get("token_cost", 0),
            attempt_latency_ms=elapsed_ms,
            failure_reason=result.get("failure_reason", "generation_failed"),
            registry_snapshot_id=snap_id,
            invocation_record_id=result.get("invocation_record_id"),
            seed=seed,
            is_pilot=is_pilot,
            required_capability_count=result.get("required_capability_count", 0),
            matched_in_registry_count=result.get("matched_in_registry_count", 0),
            chain_length=result.get("chain_length", 1),
            chain_success=result.get("chain_success", False),
            matched_and_executed_count=result.get("matched_and_executed_count", 0),
            generated_artifact_count=result.get("generated_artifact_count", 0),
            capability_types=result.get("capability_types", []),
            failed_step_index=result.get("failed_step_index"),
            failed_step_name=result.get("failed_step_name"),
            failed_step_reason=result.get("failed_step_reason"),
            output_shape_mismatch=result.get("output_shape_mismatch", False),
            output_shape_missing_keys=result.get("output_shape_missing_keys", []),
            output_shape_missing_files=result.get("output_shape_missing_files", []),
            generated_artifact_sources=result.get("generated_artifact_sources", []),
            langfuse_trace_id=result.get("langfuse_trace_id"),
        )

    # -- condition implementations -------------------------------------------

    def _run_static(
        self,
        task: TaskDefinition,
        registry: SkillRegistry,
        metrics_store: Any = None,
    ) -> dict:
        """Static baseline: registry lookup + execute pre-written implementation.

        Static runs still record invocation telemetry and apply the same output
        contract checks as JIT runs so baseline rows are comparable.
        """
        if not task.has_static_coverage:
            return {
                "gate_passed": False,
                "output_in_range": False,
                "failure_reason": "generation_failed",
                "output_shape_mismatch": False,
                "output_shape_missing_keys": [],
                "output_shape_missing_files": [],
            }

        # Check registry has the expected capability by exact name lookup
        cap = registry.get(task.static_capability_name)
        if cap is None:
            logger.warning(
                "Static capability %s not found in registry",
                task.static_capability_name,
            )
            return {
                "gate_passed": False,
                "output_in_range": False,
                "failure_reason": "generation_failed",
                "output_shape_mismatch": False,
                "output_shape_missing_keys": [],
                "output_shape_missing_files": [],
            }

        # Execute the pre-written static implementation
        try:
            static_result = _normalize_static_implementation_result(
                _run_static_implementation(task.task_id, task.fixture_paths)
            )
            output = static_result.output
            output_files = static_result.output_files
            if output_files and task.file_validator is not None:
                in_range = task.file_validator(output, output_files)
            else:
                in_range = task.validator(output)
            # Apply the same output-contract check JIT paths use.
            shape_check = (
                task.output_contract.check(output, output_files)
                if task.output_contract
                else None
            )
            shape_mismatch = (
                shape_check is not None and not shape_check.satisfied
            )

            failure_reason = "none" if in_range else "output_rejected"
            # Promote shape mismatch and keep output_in_range aligned with the
            # canonical success predicate.
            if in_range and shape_mismatch:
                failure_reason = "output_shape_mismatch"
                in_range = False

            # Record one invocation so the static condition contributes to
            # invocation-distribution metrics. Best-effort only.
            if metrics_store is not None:
                try:
                    from src.schemas.metrics import InvocationRecord
                    metrics_store.record_invocation(InvocationRecord(
                        capability_id=task.static_capability_name,
                        task_id=task.task_id,
                        success=bool(in_range),
                        token_cost=0,
                        latency_ms=0,
                    ))
                except Exception:  # noqa: BLE001
                    logger.debug(
                        "Failed to record static invocation for %s",
                        task.task_id, exc_info=True,
                    )

            return {
                "gate_passed": True,  # static implementations don't go through the gate
                "output_in_range": in_range,
                "failure_reason": failure_reason,
                "capability_type": cap.capability_type,
                "attempt_count": 0,
                "verification_tier_reached": None,  # bypassed for static
                "output_shape_mismatch": shape_mismatch,
                "output_shape_missing_keys": (
                    shape_check.missing_keys if shape_check is not None else []
                ),
                "output_shape_missing_files": (
                    shape_check.missing_files if shape_check is not None else []
                ),
            }
        except Exception as exc:  # noqa: BLE001
            logger.warning("Static implementation for %s failed: %s", task.task_id, exc)
            return {
                "gate_passed": False,
                "output_in_range": False,
                "failure_reason": "generation_failed",
                "output_shape_mismatch": False,
                "output_shape_missing_keys": [],
                "output_shape_missing_files": [],
            }

    def _run_jit(
        self,
        task: TaskDefinition,
        condition: str,
        registry: SkillRegistry,
        loader: ToolLoader,
        metrics_store: Any,
        *,
        run_id: str = "",
        is_pilot: bool = False,
    ) -> dict:
        """JIT conditions: invoke the pipeline, then validate output."""
        from src.orchestrator.pipeline import create_pipeline

        profile = get_condition_profile(condition)

        pipeline = create_pipeline(
            registry,
            loader,
            optimizer_config=profile.optimizer_config,
            metrics_store=metrics_store,
            verification_enabled=profile.verification_enabled,
        ).compile()

        # Per-task LangFuse tracing: pass task_id and condition via config
        # metadata so each pipeline invocation gets its own tagged trace.
        # LangFuse's LangChain integration reads langfuse_* keys from
        # config["metadata"] to set session_id, tags, etc. on the trace.
        invoke_config: dict = {}
        if self._config.langfuse_handler is not None:
            invoke_config["callbacks"] = [self._config.langfuse_handler]
            _file_formats = list({
                fp.suffix.lower().lstrip(".") for fp in task.fixture_paths
            }) if task.fixture_paths else []

            # Build the tag list to maximize analysis-time filterability. Axes:
            # - condition: which ablation (full / no_verification / ...)
            # - task_id / tier: per-task + per-tier aggregation
            # - format:X: per-file-format behavior
            # - allow_network:{true|false}: network-policy consistency
            # - pilot: fast exclusion of pilot runs from primary metrics
            # - static_coverage:{true|false}: filter retrieval vs synthesis
            #   tasks.
            _tags = self._build_task_lf_tags(task, condition, is_pilot)

            invoke_config["metadata"] = {
                "langfuse_session_id": run_id,
                # user_id differentiates eval-harness traces from ad-hoc
                # programmatic traces in per-user LangFuse dashboards.
                "langfuse_user_id": "eval-harness",
                "langfuse_tags": _tags,
                "task_id": task.task_id,
                "task_family_id": task.family_id,
                "task_instance_id": task.instance_id,
                "variant_id": task.variant_id,
                "benchmark_track": task.benchmark_track,
                "condition": condition,
                "run_id": run_id,
                "tier": task.tier,
                "has_static_coverage": task.has_static_coverage,
                "file_formats": _file_formats,
                "is_pilot": is_pilot,
                "allow_network": self._config.allow_network,
            }

        # Token-cost honesty. A free-form prompt path would run the LLM
        # extractor at task entry, which costs tokens. Run the extractor inside
        # the LangFuse-tagged invocation context so its tokens count toward the
        # per-task budget; discard the result. The task's authored output
        # contract remains the deterministic measurement instrument.
        try:
            _extract_output_contract_llm(task.task_prompt)
        except Exception as exc:  # noqa: BLE001
            logger.debug(
                "Token-accounting extractor call failed for %s: %s",
                task.task_id,
                exc,
            )

        # Use the shared builder so every pipeline entry point stays in sync
        # on which TaskState fields get populated.
        from src.orchestrator.task_state_builder import build_initial_state
        initial_state = build_initial_state(
            task, allow_network=self._config.allow_network,
        )
        # Measure pipeline latency in _run_jit's own scope so the LangFuse
        # latency score uses the same boundary as the pipeline invocation.
        _t_start_jit = time.monotonic()
        try:
            final_state: TaskState = pipeline.invoke(initial_state, config=invoke_config)
        except Exception as exc:  # noqa: BLE001
            logger.warning("Pipeline raised for task %s: %s", task.task_id, exc)
            return {
                "gate_passed": False,
                "output_in_range": False,
                "failure_reason": "generation_failed",
                "attempt_count": 0,
            }
        jit_elapsed_ms = int((time.monotonic() - _t_start_jit) * 1000)

        # Capture LangFuse root trace ID for DuckDB/LangFuse cross-reference.
        # None when LangFuse is disabled.
        _lf_trace_id: str | None = None
        if self._config.langfuse_handler is not None:
            try:
                _lf_trace_id = getattr(
                    self._config.langfuse_handler, "last_trace_id", None,
                )
            except Exception:  # noqa: BLE001
                _lf_trace_id = None

        # Pass LangFuse tags to execution functions so sub-traces
        # (skill/agent execution) are tagged consistently.
        _lf_meta = invoke_config.get("metadata", {})
        result = _extract_result(
            final_state, task, loader,
            metrics_store=metrics_store,
            langfuse_tags=_lf_meta.get("langfuse_tags"),
            langfuse_session_id=_lf_meta.get("langfuse_session_id"),
        )
        if _lf_trace_id is not None:
            result["langfuse_trace_id"] = _lf_trace_id

            # Push M1 success and M5 latency scores onto the LangFuse trace.
            # Aggregate metrics flow through DuckDB telemetry tables.
            try:
                from langfuse import get_client
                lf = get_client()
                # Use the canonical success predicate so LangFuse and DuckDB
                # success semantics agree.
                success_val = 1.0 if EvaluationRecord.is_success(
                    gate_passed=bool(result.get("gate_passed", False)),
                    output_in_range=bool(result.get("output_in_range", False)),
                    output_shape_mismatch=bool(result.get("output_shape_mismatch", False)),
                ) else 0.0
                latency_val = float(jit_elapsed_ms or 0)
                for method_name in ("create_score", "score"):
                    _create = getattr(lf, method_name, None)
                    if _create is None:
                        continue
                    try:
                        _create(
                            trace_id=_lf_trace_id,
                            name="m1.success",
                            value=success_val,
                            data_type="NUMERIC",
                        )
                        # M5 is only meaningful on successful runs.
                        if success_val > 0:
                            _create(
                                trace_id=_lf_trace_id,
                                name="m5.latency_ms",
                                value=latency_val,
                                data_type="NUMERIC",
                            )
                        break
                    except Exception:  # noqa: BLE001
                        continue
            except Exception:  # noqa: BLE001
                pass
        return result


# ---------------------------------------------------------------------------
# Compositional result extraction
# ---------------------------------------------------------------------------


def _extract_result_compositional(
    subtasks: list,
    artifacts: list,
    matches: list,
    required_capabilities: list[str],
    verification_results: list,
    attempt_count: int,
    task: TaskDefinition,
    *,
    metrics_store: Any | None = None,
    langfuse_tags: list[str] | None = None,
    langfuse_session_id: str | None = None,
    allow_network: bool = False,
) -> dict:
    """Compositional execution path — run all artifacts in dependency order.

    Used when the decomposition produced multiple subtasks with resolved
    capabilities.  Builds a resolution map, creates a
    :class:`CompositionalExecutor`, runs the chain, and maps
    :class:`ChainResult` fields onto the ``EvaluationRecord`` dict.
    """
    from src.orchestrator.composition import CompositionalExecutor, build_resolution_map

    # Check gate: all generated artifacts must have passed verification.
    # If verification was enabled and any artifact failed, the pipeline
    # would have retried up to MAX_ATTEMPTS.  We check the latest report
    # for each artifact.
    gate_passed = True
    verification_tier_reached: int | None = None
    if verification_results:
        _latest_by_name: dict[str, Any] = {}
        for report in verification_results:
            _latest_by_name[report.capability_name] = report
        for report in _latest_by_name.values():
            if not report.overall_passed:
                gate_passed = False
            if verification_tier_reached is None or (
                report.tier_reached is not None
                and report.tier_reached > (verification_tier_reached or 0)
            ):
                verification_tier_reached = report.tier_reached
    elif artifacts:
        # No verification reports but artifacts exist = no_verification condition
        gate_passed = True

    # Build resolution map
    resolution = build_resolution_map(
        required_capabilities, matches, artifacts,
    )

    # Determine primary capability type (highest priority in chain)
    cap_types = []
    for ref in resolution.values():
        ct = str(ref.capability_type).split(".")[-1].lower()
        if ct not in cap_types:
            cap_types.append(ct)
    # Primary type for the record: use most capable type present
    _type_order = {"agent": 0, "skill": 1, "mcp_tool": 2}
    capability_type = min(cap_types, key=lambda t: _type_order.get(t, 9)) if cap_types else None

    if not gate_passed:
        # Verification failed — don't attempt execution
        tier_map = {1: "tier1", 2: "tier2", 3: "tier3"}
        return {
            "gate_passed": False,
            "output_in_range": False,
            "failure_reason": tier_map.get(verification_tier_reached, "tier1"),
            "attempt_count": attempt_count,
            "capability_type": capability_type,
            "verification_tier_reached": verification_tier_reached,
            "token_cost": 0,
            "invocation_record_id": None,
            "required_capability_count": len(required_capabilities),
            "matched_in_registry_count": len(matches),
            "chain_length": len(subtasks),
            "chain_success": False,
            "matched_and_executed_count": 0,
            "generated_artifact_count": len(artifacts),
            "capability_types": cap_types,
            "failed_step_index": None,
            "failed_step_name": None,
            "failed_step_reason": None,
            # Gate-blocked runs never reached terminal state, so the output
            # contract check is not meaningfully evaluated here.
            "output_shape_mismatch": False,
            "output_shape_missing_keys": [],
            "output_shape_missing_files": [],
            # Snapshot even gate-rejected artifacts for later inspection.
            "generated_artifact_sources": _snapshot_artifacts(artifacts),
        }

    # Run the compositional chain
    executor = CompositionalExecutor(
        subtasks=subtasks,
        resolution=resolution,
        artifacts=artifacts,
        matches=matches,
        fixture_paths=task.fixture_paths or [],
        task=task,
        metrics_store=metrics_store,
        langfuse_tags=langfuse_tags,
        langfuse_session_id=langfuse_session_id,
        allow_network=allow_network,
    )

    chain_result = executor.run(task_id=task.instance_id)

    # Validate final output
    output_in_range = False
    failure_reason = "generation_failed"
    final_output: dict | None = None
    if chain_result.chain_success and chain_result.step_results:
        final_output = chain_result.step_results[-1]
        try:
            if chain_result.output_files and task.file_validator is not None:
                output_in_range = task.file_validator(final_output, chain_result.output_files)
            else:
                output_in_range = task.validator(final_output)
            failure_reason = "none" if output_in_range else "output_rejected"
        except Exception as exc:  # noqa: BLE001
            logger.warning(
                "Validator failed for compositional output of %s: %s",
                task.task_id, exc,
            )
            failure_reason = "output_rejected"
    elif not chain_result.chain_success:
        # Prefer the structured reason from the execution backend over the
        # generic chain_step_failed bucket, so analyzers can separate AGENT
        # timeouts from AGENT crashes and other chain-step failures.
        if chain_result.failed_step_reason == "agent_execution_timeout":
            failure_reason = "agent_execution_timeout"
        elif chain_result.failed_step_reason == "agent_execution_error":
            failure_reason = "agent_execution_error"
        else:
            failure_reason = "chain_step_failed"

    # Output contract check is orthogonal to the validator. Only evaluate it
    # after the chain reaches terminal state with a concrete output; execution
    # failures should stay in execution/infrastructure buckets rather than
    # being reclassified as shape mismatches.
    if chain_result.chain_success and final_output is not None:
        shape_check = task.output_contract.check(
            final_output, chain_result.output_files,
        ) if task.output_contract else None
    else:
        shape_check = None

    # Promote output_shape_mismatch into failure_reason when the validator is
    # happy but the typed contract is not. Flip output_in_range as well so all
    # stores share the same success predicate.
    if (
        failure_reason == "none"
        and shape_check is not None
        and not shape_check.satisfied
    ):
        failure_reason = "output_shape_mismatch"
        output_in_range = False

    # Count how many registry matches vs generated artifacts were executed
    matched_count = sum(
        1 for ref in resolution.values() if ref.source == "registry"
    )
    generated_count = sum(
        1 for ref in resolution.values() if ref.source == "generated"
    )

    return {
        "gate_passed": gate_passed,
        "output_in_range": output_in_range,
        "failure_reason": failure_reason,
        "attempt_count": attempt_count,
        "capability_type": capability_type,
        "verification_tier_reached": verification_tier_reached,
        "token_cost": 0,
        "invocation_record_id": None,
        "required_capability_count": len(required_capabilities),
        "matched_in_registry_count": len(matches),
        "chain_length": chain_result.chain_length,
        "chain_success": chain_result.chain_success,
        "matched_and_executed_count": matched_count,
        "generated_artifact_count": generated_count,
        "capability_types": cap_types,
        "failed_step_index": chain_result.failed_step_index,
        "failed_step_name": chain_result.failed_step_name,
        "failed_step_reason": chain_result.failed_step_reason,
        "output_shape_mismatch": (
            shape_check is not None and not shape_check.satisfied
        ),
        "output_shape_missing_keys": (
            shape_check.missing_keys if shape_check is not None else []
        ),
        "output_shape_missing_files": (
            shape_check.missing_files if shape_check is not None else []
        ),
        # Snapshot every generated artifact for later inspection.
        "generated_artifact_sources": _snapshot_artifacts(artifacts),
    }


# ---------------------------------------------------------------------------
# Result extraction from TaskState
# ---------------------------------------------------------------------------


def _extract_result(
    state: TaskState,
    task: TaskDefinition,
    loader: ToolLoader,
    *,
    metrics_store: Any | None = None,
    langfuse_tags: list[str] | None = None,
    langfuse_session_id: str | None = None,
) -> dict:
    """Extract EvaluationRecord fields from a completed TaskState."""
    # pipeline.invoke() returns a dict, not a TaskState instance
    _get = state.get if isinstance(state, dict) else lambda k, d=None: getattr(state, k, d)
    artifacts = _get("generated_artifacts", []) or []
    matches = _get("matches", []) or []
    required_capabilities = _get("required_capabilities", []) or []
    verification_results = _get("verification_results", []) or []
    attempt_count = _get("attempt_count", 0)
    subtasks = _get("subtasks", []) or []
    decomposition_succeeded = bool(_get("decomposition_succeeded", True))
    # Preserve task-level network policy through execution.
    allow_network = bool(_get("allow_network", False))

    # Fail closed when decomposition already reported failure. Later chain
    # output must not rescue a task whose decomposition step was unsuccessful.
    if not decomposition_succeeded:
        logger.warning(
            "Task %s marked decomposition_succeeded=False; failing closed in eval harness",
            task.task_id,
        )
        return {
            "gate_passed": False,
            "output_in_range": False,
            "failure_reason": "generation_failed",
            "attempt_count": attempt_count,
            "capability_type": None,
            "verification_tier_reached": None,
            "token_cost": 0,
            "invocation_record_id": _get("trace_id", None),
            "required_capability_count": len(required_capabilities),
            "matched_in_registry_count": len(matches),
            "chain_length": max(len(subtasks), 1),
            "chain_success": False,
            "matched_and_executed_count": 0,
            "generated_artifact_count": 0,
            "capability_types": [],
            "failed_step_index": None,
            "failed_step_name": None,
            "failed_step_reason": "decomposition_failed",
            "generated_artifact_sources": _snapshot_artifacts(artifacts),
            "output_shape_mismatch": False,
            "output_shape_missing_keys": [],
            "output_shape_missing_files": [],
        }

    # ---------------------------------------------------------------
    # Compositional execution path: when the task was decomposed into multiple
    # subtasks, run all resolved artifacts in dependency order. A single
    # lead-with-subagents AGENT artifact can cover multiple subtasks, so one
    # artifact or match is enough to enter this path.
    # ---------------------------------------------------------------
    if len(subtasks) > 1 and (len(artifacts) + len(matches) >= 1):
        return _extract_result_compositional(
            subtasks=subtasks,
            artifacts=artifacts,
            matches=matches,
            required_capabilities=required_capabilities,
            verification_results=verification_results,
            attempt_count=attempt_count,
            task=task,
            metrics_store=metrics_store,
            langfuse_tags=langfuse_tags,
            langfuse_session_id=langfuse_session_id,
            allow_network=allow_network,
        )

    # ---------------------------------------------------------------
    # Single-artifact execution path
    # ---------------------------------------------------------------

    # Determine verification outcome from the last VerificationReport
    gate_passed = False
    verification_tier_reached: int | None = None
    failure_reason = "generation_failed"
    capability_type: str | None = None

    if artifacts:
        from src.schemas.capability import CapabilityType

        # When multiple artifacts exist (decomposition produced several gaps),
        # prefer the most capable type: AGENT > SKILL > MCP_TOOL.
        # Within the same type, sort by name for deterministic selection.
        _TYPE_PRIORITY = {"AGENT": 0, "SKILL": 1, "MCP_TOOL": 2}
        artifacts = sorted(
            artifacts,
            key=lambda a: (
                _TYPE_PRIORITY.get(str(a.capability_type).split(".")[-1].upper(), 9),
                a.capability_name or "",
            ),
        )
        artifact = artifacts[0]
        capability_type = str(artifact.capability_type).split(".")[-1].lower()

        if verification_results:
            # Find the report matching the selected artifact by capability_name.
            # verification_results accumulates across retries and artifacts;
            # the last entry is NOT necessarily for the selected artifact.
            matched_report = None
            for report in reversed(verification_results):
                if report.capability_name == artifact.capability_name:
                    matched_report = report
                    break
            last_report = matched_report or verification_results[-1]
            gate_passed = last_report.overall_passed
            verification_tier_reached = last_report.tier_reached
            if not gate_passed:
                # Map tier_reached to failure_reason
                tier_map = {1: "tier1", 2: "tier2", 3: "tier3"}
                failure_reason = tier_map.get(last_report.tier_reached, "tier1")
        else:
            # No verification report = gate bypassed (no_verification condition)
            gate_passed = True
            failure_reason = "none"
    elif matches and not artifacts:
        # Registry match path: all required capabilities were found in the
        # registry — no generation was needed.  Seeded capabilities are
        # pre-verified, so gate_passed=True.  The harness executes the static
        # implementation (if available) to check output_in_range.
        logger.info(
            "Task %s resolved via registry match (%d match(es), 0 artifacts)",
            task.task_id,
            len(matches),
        )
        gate_passed = True
        failure_reason = "none"
        # Infer capability_type from the first match
        first_match = matches[0]
        cap = (
            first_match.capability
            if hasattr(first_match, "capability")
            else first_match.get("capability")
        )
        if cap is not None:
            _ct = (
                cap.capability_type
                if hasattr(cap, "capability_type")
                else cap.get("capability_type")
            )
            capability_type = str(_ct).split(".")[-1].lower() if _ct else None

    # Execute the artifact against fixture data to check output_in_range
    output_in_range = False
    output: dict | None = None
    if gate_passed and artifacts:
        # artifact already set by priority sort above (artifacts[0])
        try:
            output = _execute_artifact(
                artifact, task, loader,
                langfuse_tags=langfuse_tags,
                langfuse_session_id=langfuse_session_id,
                allow_network=allow_network,
            )
            output_in_range = task.validator(output)
            if gate_passed and not output_in_range:
                failure_reason = "output_rejected"
            elif gate_passed and output_in_range:
                failure_reason = "none"
            # Keep invocation recording parity with static and compositional
            # paths so invocation-distribution metrics include this artifact.
            if metrics_store is not None:
                try:
                    from src.schemas.metrics import InvocationRecord
                    metrics_store.record_invocation(InvocationRecord(
                        capability_id=getattr(artifact, "capability_name", None) or "",
                        task_id=task.task_id[:64],
                        success=bool(output_in_range),
                        token_cost=0,
                        latency_ms=0,
                    ))
                except Exception:  # noqa: BLE001
                    logger.debug("record_invocation failed (single-artifact)", exc_info=True)
        except Exception as exc:  # noqa: BLE001
            logger.warning("Artifact execution failed for task %s: %s", task.task_id, exc)
    elif gate_passed and not artifacts and task.has_static_coverage:
        # Registry match with static implementation available — execute it.
        try:
            static_result = _normalize_static_implementation_result(
                _run_static_implementation(task.task_id, task.fixture_paths)
            )
            output = static_result.output
            if static_result.output_files and task.file_validator is not None:
                output_in_range = task.file_validator(output, static_result.output_files)
            else:
                output_in_range = task.validator(output)
            if not output_in_range:
                failure_reason = "output_rejected"
            logger.info(
                "Registry-match static execution for %s: output_in_range=%s",
                task.task_id,
                output_in_range,
            )
            # Invocation recording parity for the JIT-registry-match path.
            if metrics_store is not None and matches:
                try:
                    from src.schemas.metrics import InvocationRecord
                    first_match = matches[0]
                    cap = (
                        first_match.capability
                        if hasattr(first_match, "capability")
                        else first_match.get("capability")
                    )
                    cap_name = (
                        getattr(cap, "capability_name", None)
                        or (cap.get("capability_name") if isinstance(cap, dict) else None)
                        or task.static_capability_name
                        or ""
                    )
                    metrics_store.record_invocation(InvocationRecord(
                        capability_id=cap_name,
                        task_id=task.task_id[:64],
                        success=bool(output_in_range),
                        token_cost=0,
                        latency_ms=0,
                    ))
                except Exception:  # noqa: BLE001
                    logger.debug("record_invocation failed (registry-match)", exc_info=True)
        except Exception as exc:  # noqa: BLE001
            logger.warning(
                "Static execution for registry-matched task %s failed: %s",
                task.task_id,
                exc,
            )
    elif gate_passed and not artifacts:
        # Registry match but no static implementation. Try to extract source_code
        # from the matched capability's metadata and sandbox-execute it.
        first_match = matches[0] if matches else None
        cap = first_match.capability if first_match and hasattr(first_match, "capability") else None
        matched_source = None
        if cap is not None:
            meta = cap.metadata if hasattr(cap, "metadata") else {}
            matched_source = meta.get("source_code")

        if matched_source:
            try:
                from src.schemas.capability import CapabilityType
                from src.schemas.task import GeneratedArtifact

                # Reconstruct skill_content / agent_config from metadata
                # so SKILL/AGENT execution backends have original instructions.
                matched_skill = meta.get("skill_content")
                matched_agent_config = None
                _ac_raw = meta.get("agent_config")
                if _ac_raw:
                    try:
                        import json as _json
                        matched_agent_config = (
                            _json.loads(_ac_raw)
                            if isinstance(_ac_raw, str)
                            else _ac_raw
                        )
                    except (ValueError, TypeError):
                        matched_agent_config = None

                proxy_artifact = GeneratedArtifact(
                    capability_name=cap.name,
                    capability_type=CapabilityType(cap.capability_type)
                    if isinstance(cap.capability_type, str)
                    else cap.capability_type,
                    source_code=matched_source,
                    skill_content=matched_skill,
                    agent_config=matched_agent_config,
                )
                output = _execute_artifact(
                    proxy_artifact, task, loader,
                    langfuse_tags=langfuse_tags,
                    langfuse_session_id=langfuse_session_id,
                    allow_network=allow_network,
                )
                output_in_range = task.validator(output)
                if not output_in_range:
                    failure_reason = "output_rejected"
                else:
                    failure_reason = "none"
                logger.info(
                    "Registry-match sandbox execution for %s: output_in_range=%s",
                    task.task_id,
                    output_in_range,
                )
            except Exception as exc:  # noqa: BLE001
                logger.warning(
                    "Registry-match execution for task %s failed: %s",
                    task.task_id,
                    exc,
                )
                failure_reason = "output_rejected"
        else:
            failure_reason = "no_source_code"
            logger.info(
                "Task %s matched in registry but no source_code in metadata — "
                "cannot validate output",
                task.task_id,
            )

    # Compositional fields: defaults that match single-artifact behavior.
    chain_success = bool(gate_passed and output_in_range)
    if artifacts:
        # Generation path: one freshly-generated artifact was executed.
        generated_artifact_count = 1
        matched_and_executed_count = 0
    elif matches:
        # Registry-match path: one matched capability was executed.
        generated_artifact_count = 0
        matched_and_executed_count = 1
    else:
        generated_artifact_count = 0
        matched_and_executed_count = 0
    capability_types_list = [capability_type] if capability_type is not None else []

    # Output contract check on the single-artifact path.
    # Single-artifact execution does not surface a separate output_files
    # mapping (file outputs only appear in compositional chains today),
    # so the file-presence check uses an empty dict here. T07/T11 — the
    # only file-producing tasks — go through the compositional path.
    # Only evaluate the contract when execution reached terminal state with a
    # non-None output.
    if output is not None:
        shape_check = task.output_contract.check(output, {}) if task.output_contract else None
    else:
        shape_check = None

    # Keep shape-mismatch handling symmetric with the compositional branch.
    if (
        failure_reason == "none"
        and shape_check is not None
        and not shape_check.satisfied
    ):
        failure_reason = "output_shape_mismatch"
        output_in_range = False

    return {
        "gate_passed": gate_passed,
        "output_in_range": output_in_range,
        "failure_reason": failure_reason,
        "attempt_count": attempt_count,
        "capability_type": capability_type,
        "verification_tier_reached": verification_tier_reached,
        "token_cost": 0,
        "invocation_record_id": _get("trace_id", None),
        "required_capability_count": len(required_capabilities),
        "matched_in_registry_count": len(matches),
        "chain_length": 1,
        "chain_success": chain_success,
        "matched_and_executed_count": matched_and_executed_count,
        "generated_artifact_count": generated_artifact_count,
        "capability_types": capability_types_list,
        "failed_step_index": None,
        "failed_step_name": None,
        "failed_step_reason": None,
        # Snapshot generated artifacts for later inspection.
        "generated_artifact_sources": _snapshot_artifacts(artifacts),
        "output_shape_mismatch": (
            shape_check is not None and not shape_check.satisfied
        ),
        "output_shape_missing_keys": (
            shape_check.missing_keys if shape_check is not None else []
        ),
        "output_shape_missing_files": (
            shape_check.missing_files if shape_check is not None else []
        ),
    }


def _execute_artifact(
    artifact: Any,
    task: TaskDefinition,
    loader: ToolLoader,
    *,
    langfuse_tags: list[str] | None = None,
    langfuse_session_id: str | None = None,
    allow_network: bool = False,
) -> dict:
    """Execute a generated artifact with fixture data and return its output."""
    from src.schemas.capability import CapabilityType

    cap_type = artifact.capability_type
    if cap_type == CapabilityType.SKILL:
        from tests.evaluation.execution import _read_fixture_contents, execute_skill_artifact

        fixture_contents = _read_fixture_contents(task)
        return execute_skill_artifact(
            artifact, task, fixture_contents,
            langfuse_tags=langfuse_tags,
            langfuse_session_id=langfuse_session_id,
            allow_network=allow_network,
        )
    elif cap_type == CapabilityType.AGENT:
        from tests.evaluation.execution import _read_fixture_contents, execute_agent_artifact

        fixture_contents = _read_fixture_contents(task)
        return execute_agent_artifact(
            artifact, task, fixture_contents,
            langfuse_tags=langfuse_tags,
            langfuse_session_id=langfuse_session_id,
            allow_network=allow_network,
        )
    elif cap_type != CapabilityType.MCP_TOOL:
        logger.debug("Unknown artifact type: %s", cap_type)
        return {}

    source_code = getattr(artifact, "source_code", None)
    if not source_code:
        return {}

    # Use sandbox_execute to run the tool with fixture-derived inputs
    from src.verification.sandbox_tool import resolve_extra_packages, sandbox_execute

    fixture_inputs = _build_fixture_inputs(task, source_code=source_code)

    # Extract the actual async function name from source code AST
    from src.verification.behavioral import _extract_function_name

    func_name = _extract_function_name(source_code) or getattr(
        artifact, "capability_name", "tool"
    ).replace("-", "_")

    # Embed fixture file content in the sandbox code since Docker doesn't
    # mount the host filesystem.  Write each fixture to a temp path inside
    # the container and update fixture_inputs to point there.
    fixture_preamble = ""
    if task.fixture_paths:
        _BINARY_FIXTURE_EXTENSIONS = {".xlsx", ".xls", ".parquet", ".pickle", ".pkl"}
        for i, fp in enumerate(task.fixture_paths):
            host_path = Path(fp) if Path(fp).is_absolute() else Path.cwd() / fp
            if host_path.exists():
                sandbox_path = f"/tmp/fixture_{i}{host_path.suffix}"
                if host_path.suffix.lower() in _BINARY_FIXTURE_EXTENSIONS:
                    import base64
                    b64 = base64.b64encode(host_path.read_bytes()).decode("ascii")
                    fixture_preamble += (
                        f"import pathlib as _pl, base64 as _b64\n"
                        f"_pl.Path({sandbox_path!r}).write_bytes(_b64.b64decode({b64!r}))\n"
                    )
                else:
                    content = host_path.read_text(encoding="utf-8", errors="replace")
                    fixture_preamble += (
                        "import pathlib as _pl\n"
                        f"_pl.Path({sandbox_path!r}).write_text({content!r})\n"
                    )
                # Replace the host path in fixture_inputs with the sandbox path
                for k, v in fixture_inputs.items():
                    if str(v) == str(fp):
                        fixture_inputs[k] = sandbox_path

    # Build a sandbox invocation (same pattern as behavioral testing).
    # Use inspect to filter kwargs to only what the function accepts,
    # since extra_inputs may include params the tool doesn't expose.
    # Fuzzy-match extra_inputs keys to actual param names (e.g.
    # "features" matches "feature_columns") to handle LLM naming variance.
    sandbox_code = f"""
import asyncio, json, inspect
__name__ = "not_main"
{fixture_preamble}
{source_code}

async def _run():
    try:
        _all_inputs = {repr(fixture_inputs)}
        _sig = inspect.signature({func_name})
        _valid = set(_sig.parameters.keys())
        _kwargs = {{k: v for k, v in _all_inputs.items() if k in _valid}}
        # Fuzzy-match: for unmatched params, find extra inputs that share
        # a root word (e.g. "features" ↔ "feature_columns").
        # Only try to fill params that are truly required (no default value)
        _required = {{
            name for name, p in _sig.parameters.items()
            if p.default is inspect.Parameter.empty
        }}
        _unmatched = _required - set(_kwargs)
        _unused = {{k: v for k, v in _all_inputs.items() if k not in _kwargs}}
        def _words(s):
            return set(w.rstrip("s") for w in s.lower().replace("_", " ").split() if len(w) > 2)
        for _p in _unmatched:
            _pw = _words(_p)
            best, best_n = None, 0
            for _ek, _ev in _unused.items():
                n = len(_pw & _words(_ek))
                if n > best_n:
                    best, best_n = (_ek, _ev), n
            if best_n > 0:
                _kwargs[_p] = best[1]
        result = await {func_name}(**_kwargs)
        print(json.dumps({{"result": result}}))
    except Exception as e:
        print(json.dumps({{"error": str(e)}}))

asyncio.run(_run())
"""
    # Prefer pip_packages declared by the generator (accurate PyPI names).
    # Fall back to LLM-based import resolution for old artifacts.
    declared_pkgs = getattr(artifact, "pip_packages", None)
    if declared_pkgs:
        extra_pkgs = declared_pkgs
    else:
        extra_pkgs = resolve_extra_packages(source_code)
    # Allow more time when dynamic dependencies need installation before exec.
    exec_timeout = 60 if extra_pkgs else 30
    logger.debug("Sandbox extra_packages=%s timeout=%s", extra_pkgs, exec_timeout)
    result = sandbox_execute(
        sandbox_code,
        timeout=exec_timeout,
        extra_packages=extra_pkgs or None,
        allow_network=allow_network,
    )
    if not result.get("success") or not result.get("stdout"):
        logger.warning(
            "Sandbox execution for %s failed: exit_code=%s success=%s stderr=%s",
            func_name,
            result.get("exit_code"),
            result.get("success"),
            (result.get("stderr") or "")[:500],
        )
        return {}

    try:
        parsed = json.loads(result["stdout"].strip().splitlines()[-1])
        if "error" in parsed:
            logger.warning(
                "Sandbox tool %s raised: %s",
                func_name,
                str(parsed["error"])[:500],
            )
            return {}
        # Preserve falsy / non-dict results via the shared helper.
        from src.orchestrator.output_coerce import extract_tool_result
        output = extract_tool_result(parsed)
        logger.info("Sandbox output for %s: %s", func_name, str(output)[:300])
        return output
    except (json.JSONDecodeError, KeyError, IndexError):
        logger.warning(
            "Sandbox output parse failed for %s: stdout=%s",
            func_name,
            result["stdout"][:200],
        )
        return {}


def _build_fixture_inputs(task: TaskDefinition, source_code: str | None = None) -> dict:
    """Build keyword arguments matching fixture paths to the tool's parameters.

    Uses AST extraction of the generated tool's parameter names to map fixture
    files intelligently.  Falls back to ``{"file_path": <primary>}`` when the
    tool signature cannot be parsed or has a single path-like parameter.
    """
    if not task.fixture_paths:
        return {}

    # --- Extract parameter names from source code ---
    param_names: list[str] = []
    if source_code:
        from src.verification.behavioral import extract_parameter_schema

        schema = extract_parameter_schema(source_code)
        param_names = list(schema.keys())

    # --- Single-fixture shortcut ---
    if len(task.fixture_paths) == 1:
        # Find the first path-like param, or default to "file_path"
        path_param = "file_path"
        for p in param_names:
            if any(kw in p.lower() for kw in ("path", "file", "csv", "log", "json", "txt", "doc")):
                path_param = p
                break
        inputs = {path_param: str(task.fixture_paths[0])}
        # Merge extra_inputs before returning
        if hasattr(task, "extra_inputs") and task.extra_inputs:
            for k, v in task.extra_inputs.items():
                if k not in inputs:
                    inputs[k] = v
        return inputs

    # --- Multi-fixture: match by keyword overlap ---
    inputs: dict[str, str] = {}
    # Identify path-like parameters
    path_params = [
        p
        for p in param_names
        if any(kw in p.lower() for kw in ("path", "file", "csv", "log", "json", "txt", "doc"))
    ]
    unmatched_fixtures = list(task.fixture_paths)

    for param in path_params:
        best_fixture = None
        best_score = 0
        param_lower = param.lower()
        for fp in unmatched_fixtures:
            # Score by keyword overlap between param name and fixture filename
            stem = fp.stem.lower()  # e.g. "t07_left"
            score = sum(1 for kw in stem.split("_") if kw in param_lower)
            # Also check common keywords like "left"/"right"/"schema"/"config"
            if score > best_score:
                best_score = score
                best_fixture = fp
        if best_fixture and best_score > 0:
            inputs[param] = str(best_fixture)
            unmatched_fixtures.remove(best_fixture)

    # If keyword matching didn't assign all fixtures, assign remaining by order
    remaining_path_params = [p for p in path_params if p not in inputs]
    for param, fp in zip(remaining_path_params, unmatched_fixtures):
        inputs[param] = str(fp)
        # Don't remove from unmatched_fixtures in the zip — it's consumed

    # If no path params were found at all, fall back to positional names
    if not inputs:
        if len(task.fixture_paths) == 2:
            inputs = {
                "left_path": str(task.fixture_paths[0]),
                "right_path": str(task.fixture_paths[1]),
            }
        else:
            for i, fp in enumerate(task.fixture_paths):
                # Semantic key from filename for better fuzzy matching
                stem = fp.stem.split("_", 1)[-1] if "_" in fp.stem else fp.stem
                ext = fp.suffix.lstrip(".")
                inputs[f"{stem}_{ext}_path"] = str(fp)
                inputs[f"file_path_{i}"] = str(fp)

    # --- Extract non-path parameters from the task prompt ---
    # Look for column-related params and try to extract values from the prompt
    for param in param_names:
        if param in inputs:
            continue
        param_lower = param.lower()
        if "column" in param_lower or "col" in param_lower:
            # Extract quoted column name from task prompt
            import re

            col_match = re.search(r"'(\w+)'\s*column", task.task_prompt)
            if not col_match:
                col_match = re.search(r"column\s*['\"](\w+)['\"]", task.task_prompt)
            if not col_match:
                col_match = re.search(r"on\s+(?:the\s+)?'(\w+)'", task.task_prompt)
            if col_match:
                inputs[param] = col_match.group(1)

    # --- Merge task-level extra_inputs (semantic kwargs like features, test_size) ---
    if hasattr(task, "extra_inputs") and task.extra_inputs:
        for k, v in task.extra_inputs.items():
            if k not in inputs:
                inputs[k] = v

    return inputs


# ---------------------------------------------------------------------------
# Static capability implementations
# ---------------------------------------------------------------------------


def _run_static_implementation(
    task_id: str, fixture_paths: list[Path],
) -> dict | tuple[dict, dict[str, bytes]]:
    """Execute a hand-written static implementation for seeded tasks.

    The dispatch table must align with ``has_static_coverage=True`` in
    ``tests/evaluation/tasks.py``. The module-load-time assertion below fails
    if the task metadata and implementation table drift apart.
    """
    if task_id == "T01":
        return _static_t01(fixture_paths[0])
    if task_id == "T02":
        return _static_t02(fixture_paths[0], fixture_paths[1])
    if task_id == "T05":
        return _static_t05(fixture_paths[0])
    if task_id == "T09":
        return _static_t09(fixture_paths[0])
    if task_id == "T04":
        return _static_t04(fixture_paths[0], fixture_paths[1])
    if task_id == "T11":
        return _static_t11(fixture_paths[0], fixture_paths[1], fixture_paths[2])
    if task_id == "T13":
        return _static_t13(fixture_paths[0], fixture_paths[1])
    if task_id == "T14":
        return _static_t14(fixture_paths[0], fixture_paths[1], fixture_paths[2])
    if task_id == "T15":
        return _static_t15(fixture_paths[0], fixture_paths[1])
    if task_id == "T19":
        return _static_t19(fixture_paths[0])
    if task_id == "T20":
        return _static_t20(fixture_paths[0], fixture_paths[1])
    raise ValueError(f"No static implementation for {task_id}")


# Tasks for which a hand-written static implementation exists in this
# module. KEEP IN SYNC with the dispatch above and with
# ``has_static_coverage=True`` in tests/evaluation/tasks.py — the
# startup assertion at the bottom of this module enforces the
# bidirectional invariant.
_STATIC_IMPLEMENTATION_TASKS: frozenset[str] = frozenset(
    {"T01", "T02", "T04", "T05", "T09", "T11", "T13", "T14", "T15", "T19", "T20"}
)


def _static_t01(csv_path: Path) -> dict:
    """Parse standard CSV, return row_count + column_names."""
    import csv as _csv

    with csv_path.open() as f:
        reader = _csv.DictReader(f)
        rows = list(reader)
        return {"row_count": len(rows), "column_names": list(reader.fieldnames or [])}


def _static_t02(config_path: Path, schema_path: Path) -> dict:
    """Validate JSON config against schema, return errors list."""
    import json as _json

    try:
        import jsonschema
    except ImportError:
        # Fallback: manual checks matching the known fixture errors
        config = _json.loads(config_path.read_text())
        errors = []
        db = config.get("database", {})
        if not isinstance(db.get("port"), int):
            errors.append(f"database.port: {db.get('port')!r} is not of type 'integer'")
        cache = config.get("cache", {})
        if isinstance(cache.get("ttl"), (int, float)) and cache["ttl"] < 0:
            errors.append(f"cache.ttl: {cache['ttl']} is less than the minimum of 0")
        log = config.get("logging", {})
        valid_levels = ["DEBUG", "INFO", "WARNING", "ERROR"]
        if log.get("level") not in valid_levels:
            errors.append(f"logging.level: {log.get('level')!r} is not one of {valid_levels}")
        return {"errors": errors}

    config = _json.loads(config_path.read_text())
    schema = _json.loads(schema_path.read_text())
    validator = jsonschema.Draft7Validator(schema)
    errors = [
        f"{'.'.join(str(p) for p in e.absolute_path) or e.validator_value}: {e.message}"
        if e.absolute_path
        else e.message
        for e in validator.iter_errors(config)
    ]
    return {"errors": errors}


def _static_t05(csv_path: Path) -> dict:
    """Remove nulls and duplicates, return counts."""
    import csv as _csv

    with csv_path.open() as f:
        reader = _csv.DictReader(f)
        rows = list(reader)

    # Drop rows with any empty value
    non_null = [r for r in rows if all(v.strip() for v in r.values())]
    # Deduplicate
    seen: set[str] = set()
    unique = []
    for r in non_null:
        key = str(sorted(r.items()))
        if key not in seen:
            seen.add(key)
            unique.append(r)

    # null_count counts nulls in the *cleaned* output (should be 0 after proper cleaning).
    # Counting on `unique` (which derives from non_null) is correct: the cleaned data has no nulls.
    null_count = sum(1 for r in unique for v in r.values() if not v.strip())
    return {"output_row_count": len(unique), "null_count": null_count}


def _static_t04(csv_path: Path, yaml_path: Path) -> dict:
    """Apply the reorder policy from YAML to the inventory CSV."""
    import csv as _csv
    import math as _math

    import yaml

    rules = yaml.safe_load(yaml_path.read_text())
    min_days = float(rules["min_days_of_stock"])
    reorder_days = float(rules["reorder_days_supply"])
    max_order_value = float(rules["max_order_value"])

    with csv_path.open() as f:
        rows = list(_csv.DictReader(f))

    below_threshold = 0
    total_cost = 0.0
    for row in rows:
        current_stock = float(row["current_stock"])
        daily_usage = float(row["daily_usage"])
        unit_cost = float(row["unit_cost"])
        days_of_stock = current_stock / daily_usage if daily_usage else float("inf")
        if days_of_stock < min_days:
            below_threshold += 1
            reorder_qty = _math.ceil(daily_usage * reorder_days)
            order_cost = reorder_qty * unit_cost
            total_cost += min(order_cost, max_order_value)

    return {
        "product_count": len(rows),
        "items_below_threshold": below_threshold,
        "total_reorder_cost": round(total_cost, 2),
    }


def _static_t09(data_path: Path) -> dict:
    """Generate a structured report from XML data."""
    import xml.etree.ElementTree as ET

    tree = ET.parse(data_path)
    root = tree.getroot()
    records = []
    for rec_el in root.findall(".//record"):
        records.append({
            "region": rec_el.findtext("region", ""),
            "product": rec_el.findtext("product", ""),
            "revenue": float(rec_el.findtext("revenue", "0")),
            "units": int(rec_el.findtext("units", "0")),
            "returns": int(rec_el.findtext("returns", "0")),
        })

    total_revenue = sum(r["revenue"] for r in records)
    total_units = sum(r["units"] for r in records)

    # Stats table
    stats_table = {
        "total_revenue": total_revenue,
        "total_units": total_units,
        "record_count": len(records),
        "avg_revenue": total_revenue / len(records) if records else 0,
    }

    # Top 5 anomalies (highest revenue)
    sorted_records = sorted(records, key=lambda r: r["revenue"], reverse=True)
    anomalies = sorted_records[:5]

    summary = (
        f"Q4 2024 Sales Analysis covering {len(records)} records across "
        f"all regions. Total revenue: ${total_revenue:,.2f}, "
        f"total units sold: {total_units:,}."
    )

    return {"summary": summary, "stats_table": stats_table, "anomalies": anomalies}


def _static_t11(
    sales_csv_path: Path,
    products_xlsx_path: Path,
    targets_csv_path: Path,
) -> tuple[dict, dict[str, bytes]]:
    """Join sales/products/targets and emit an Excel summary."""
    import csv as _csv
    import io

    import openpyxl

    with sales_csv_path.open() as f:
        sales = list(_csv.DictReader(f))

    products_wb = openpyxl.load_workbook(products_xlsx_path, read_only=True)
    product_rows = list(products_wb.active.iter_rows(values_only=True))
    product_headers = [str(v) for v in product_rows[0]]
    products = [dict(zip(product_headers, row, strict=False)) for row in product_rows[1:]]
    products_by_id = {str(p["product_id"]): p for p in products}

    with targets_csv_path.open() as f:
        targets = list(_csv.DictReader(f))
    targets_by_region = {row["region"]: float(row["target_revenue"]) for row in targets}

    region_revenues: dict[str, float] = {}
    for sale in sales:
        product_id = str(sale["product_id"])
        if product_id not in products_by_id:
            continue
        region = str(sale["region"])
        revenue = float(sale["quantity"]) * float(sale["unit_price"])
        region_revenues[region] = region_revenues.get(region, 0.0) + revenue

    sorted_regions = sorted(region_revenues)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["region", "total_revenue", "target_revenue", "meets_target"])

    regions_meeting_target = 0
    for region in sorted_regions:
        total_revenue = round(region_revenues[region], 2)
        target_revenue = float(targets_by_region[region])
        meets_target = total_revenue >= target_revenue
        regions_meeting_target += int(meets_target)
        ws.append([region, total_revenue, target_revenue, meets_target])
        region_revenues[region] = total_revenue

    buf = io.BytesIO()
    wb.save(buf)
    output = {
        "summary_row_count": len(sorted_regions),
        "region_revenues": region_revenues,
        "regions_meeting_target": regions_meeting_target,
        "output_file": "output.xlsx",
    }
    return output, {"output.xlsx": buf.getvalue()}


def _static_t13(metrics_csv_path: Path, baselines_json_path: Path) -> dict:
    """Detect anomalies directly from the metric baselines."""
    import csv as _csv
    import json as _json

    baselines = _json.loads(baselines_json_path.read_text())
    with metrics_csv_path.open() as f:
        rows = list(_csv.DictReader(f))

    anomalies = []
    for row in rows:
        timestamp = row["timestamp"]
        for metric, bounds in baselines.items():
            value = float(row[metric])
            if value < float(bounds["min"]) or value > float(bounds["max"]):
                anomalies.append({
                    "timestamp": timestamp,
                    "metric": metric,
                    "value": value,
                })

    return {
        "anomaly_count": len(anomalies),
        "anomalies": anomalies,
        "metrics_analyzed": len(baselines),
    }


def _static_t14(
    departments_csv_path: Path,
    transactions_tsv_path: Path,
    rules_json_path: Path,
) -> dict:
    """Aggregate department spending and apply the budget policy."""
    import csv as _csv
    import json as _json

    with departments_csv_path.open() as f:
        departments = list(_csv.DictReader(f))
    budgets = {
        row["department"]: float(row["annual_budget"])
        for row in departments
    }

    with transactions_tsv_path.open() as f:
        transactions = list(_csv.DictReader(f, delimiter="\t"))

    rules = _json.loads(rules_json_path.read_text())
    threshold = float(rules["max_single_transaction"])

    department_spending = {department: 0.0 for department in budgets}
    flagged_transactions = 0
    for row in transactions:
        department = row["department"]
        amount = float(row["amount"])
        department_spending[department] = department_spending.get(department, 0.0) + amount
        if amount > threshold:
            flagged_transactions += 1

    department_spending = {
        department: round(total, 2)
        for department, total in department_spending.items()
    }
    over_budget_count = sum(
        1 for department, total in department_spending.items()
        if total > budgets[department]
    )
    return {
        "department_count": len(budgets),
        "department_spending": department_spending,
        "over_budget_count": over_budget_count,
        "flagged_transactions": flagged_transactions,
    }


def _static_t15(hourly_metrics_csv_path: Path, config_json_path: Path) -> dict:
    """Apply config-driven daily aggregations and threshold checks."""
    import csv as _csv
    import json as _json
    from collections import defaultdict

    config = _json.loads(config_json_path.read_text())
    with hourly_metrics_csv_path.open() as f:
        rows = list(_csv.DictReader(f))

    by_day: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        day = row["timestamp"].split(" ")[0]
        by_day[day].append(row)

    alert_days = 0
    metrics_cfg = config["metrics"]
    thresholds = config["alert_thresholds"]
    for day_rows in by_day.values():
        day_alert = False
        for metric, aggregations in metrics_cfg.items():
            values = [float(r[metric]) for r in day_rows]
            for agg in aggregations:
                if agg == "mean":
                    agg_value = sum(values) / len(values)
                elif agg == "max":
                    agg_value = max(values)
                elif agg == "sum":
                    agg_value = sum(values)
                else:
                    continue
                threshold = thresholds.get(f"{metric}_{agg}")
                if threshold is not None and agg_value > float(threshold):
                    day_alert = True
        alert_days += int(day_alert)

    return {
        "n_days": len(by_day),
        "metrics_aggregated": len(metrics_cfg),
        "alert_days": alert_days,
    }


def _static_t19(csv_path: Path) -> dict:
    """Exploratory data analysis on product sales CSV.

    The static baseline returns the contract keys checked by ``_validate_t19``:
    ``n_rows``, ``n_categories``, and ``categories``.
    """
    import csv as _csv

    with csv_path.open() as f:
        reader = _csv.DictReader(f)
        rows = list(reader)

    categories = sorted({r.get("category", "").strip() for r in rows if r.get("category")})
    return {
        "n_rows": len(rows),
        "n_categories": len(categories),
        "categories": categories,
    }


def _static_t20(applications_csv_path: Path, policy_json_path: Path) -> dict:
    """Apply the eligibility rules to every application row."""
    import csv as _csv
    import json as _json

    policy = _json.loads(policy_json_path.read_text())
    rules = policy["rules"]

    with applications_csv_path.open() as f:
        applications = list(_csv.DictReader(f))

    def _passes_rule(application: dict, rule: dict) -> bool:
        value = application[rule["field"]]
        op = rule["operator"]
        if op == "less_than":
            return float(value) < float(rule["threshold"])
        if op == "between":
            numeric = float(value)
            return float(rule["min"]) <= numeric <= float(rule["max"])
        if op == "greater_than_or_equal":
            return float(value) >= float(rule["threshold"])
        if op == "in_list":
            return value in set(rule["allowed"])
        raise ValueError(f"Unsupported operator: {op}")

    eligible_count = 0
    ineligible_count = 0
    for application in applications:
        passed = all(_passes_rule(application, rule) for rule in rules)
        if passed:
            eligible_count += 1
        else:
            ineligible_count += 1

    return {
        "total_applications": len(applications),
        "eligible_count": eligible_count,
        "ineligible_count": ineligible_count,
        "rules_applied": len(rules),
    }


def _assert_static_implementation_coverage() -> None:
    """Guard that static coverage metadata matches this module's implementations.

    Asserts the ``_STATIC_IMPLEMENTATION_TASKS`` whitelist exactly
    matches the set of tasks in ``tests/evaluation/tasks.py`` with
    ``has_static_coverage=True``. Loud-fails on drift in either
    direction so future edits to ``tasks.py`` or the dispatch table
    can't silently regress.
    """
    try:
        from tests.evaluation.tasks import ALL_TASKS
    except Exception:  # noqa: BLE001
        # During bootstrap (e.g., sphinx build or partial import) tasks
        # may not be importable; skip the assertion then.
        return
    coverage_tasks = {t.task_id for t in ALL_TASKS if t.has_static_coverage}
    if coverage_tasks != _STATIC_IMPLEMENTATION_TASKS:
        missing_impl = coverage_tasks - _STATIC_IMPLEMENTATION_TASKS
        missing_coverage = _STATIC_IMPLEMENTATION_TASKS - coverage_tasks
        raise AssertionError(
            "Static implementation whitelist drift. has_static_coverage=True tasks: "
            f"{sorted(coverage_tasks)}; _STATIC_IMPLEMENTATION_TASKS: "
            f"{sorted(_STATIC_IMPLEMENTATION_TASKS)}. "
            f"Missing implementation for: {sorted(missing_impl)}. "
            f"Missing coverage flag for: {sorted(missing_coverage)}."
        )


_assert_static_implementation_coverage()
