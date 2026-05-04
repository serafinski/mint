"""Recurrence-task builder.

Uses the canonical task definitions as templates and swaps in:

- recurrence fixture paths
- variant-specific validators
- variant/task-instance identity
"""

from __future__ import annotations

import io
import json
from dataclasses import replace
from functools import partial
from pathlib import Path

from tests.evaluation.tasks import TaskDefinition, get_task

ROOT = Path(__file__).resolve().parents[2]
RECURRENCE_FIXTURES_DIR = ROOT / "tests" / "evaluation" / "fixtures_recurrence"
RECURRENCE_MANIFEST_PATH = RECURRENCE_FIXTURES_DIR / "manifest.json"


def _manifest_rows() -> list[dict]:
    if not RECURRENCE_MANIFEST_PATH.exists():
        raise FileNotFoundError(
            f"Recurrence manifest not found at {RECURRENCE_MANIFEST_PATH}. "
            "Run `python -m tests.evaluation.fixtures.generate_recurrence` first."
        )
    return json.loads(RECURRENCE_MANIFEST_PATH.read_text())


def _rebind_fixture_paths(task: TaskDefinition, fixture_paths: list[Path]) -> list[Path]:
    """Return recurrence fixtures in the canonical semantic order.

    Recurrence manifests store fixture files in alphabetical order, but several
    canonical tasks expect a different positional order. Rebind by exact
    filename so both prompt rewriting and any positional execution path (notably
    the static baseline) preserve canonical file roles.
    """
    by_name: dict[str, Path] = {}
    for fp in fixture_paths:
        if fp.name in by_name:
            raise ValueError(
                f"Duplicate recurrence fixture filename '{fp.name}' for task {task.task_id}"
            )
        by_name[fp.name] = fp

    ordered: list[Path] = []
    missing: list[str] = []
    for canonical_fp in task.fixture_paths:
        matched = by_name.get(canonical_fp.name)
        if matched is None:
            missing.append(canonical_fp.name)
        else:
            ordered.append(matched)

    extra = sorted(set(by_name) - {fp.name for fp in task.fixture_paths})
    if missing or extra:
        raise ValueError(
            "Recurrence fixture-role binding mismatch for "
            f"{task.task_id}: missing={missing}, extra={extra}"
        )
    return ordered


def _rewrite_prompt(task: TaskDefinition, fixture_paths: list[Path]) -> str:
    fixture_paths = _rebind_fixture_paths(task, fixture_paths)
    prompt = task.task_prompt
    for old_fp, new_fp in zip(task.fixture_paths, fixture_paths, strict=False):
        old_rel = str(old_fp.relative_to(ROOT))
        new_rel = str(new_fp.relative_to(ROOT))
        prompt = prompt.replace(old_rel, new_rel)
    return prompt


def _validate_t04_recurrence(output: dict, *, expected: dict) -> bool:
    try:
        ok_count = int(output.get("product_count", -1)) == expected["product_count"]
        actual_below = int(output.get("items_below_threshold", -1))
        ok_below = abs(actual_below - expected["items_below_threshold"]) <= 2
        actual_cost = float(output.get("total_reorder_cost", -1))
        expected_cost = expected["total_reorder_cost"]
        ok_cost = abs(actual_cost - expected_cost) / max(abs(expected_cost), 1) < 0.20
        return ok_count and ok_below and ok_cost
    except (TypeError, ValueError):
        return False


def _validate_t11_recurrence(output: dict, *, expected: dict) -> bool:
    try:
        if int(output.get("summary_row_count", -1)) != expected["summary_row_count"]:
            return False
        revenues = output.get("region_revenues", {})
        tol = expected["revenue_tolerance_pct"]
        for region, expected_rev in expected["region_revenues"].items():
            actual = revenues.get(region)
            if actual is None:
                return False
            if abs(float(actual) - expected_rev) / max(abs(expected_rev), 1.0) > (tol / 100.0):
                return False
        return int(output.get("regions_meeting_target", -1)) == expected["regions_meeting_target"]
    except (TypeError, ValueError, AttributeError):
        return False


def _validate_t11_recurrence_file(
    output: dict,
    output_files: dict[str, bytes],
    *,
    expected: dict,
) -> bool:
    if not _validate_t11_recurrence(output, expected=expected):
        return False
    xlsx_bytes = output_files.get("output.xlsx")
    if xlsx_bytes is None:
        return False
    try:
        import openpyxl
    except ImportError:
        return True
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
        if len(wb.sheetnames) == 0:
            return False
        ws = wb.active
        row_count = sum(1 for _ in ws.iter_rows(min_row=1))
        return row_count >= 5
    except Exception:
        return False


def _validate_t13_recurrence(output: dict, *, expected: dict) -> bool:
    valid_metrics = {"cpu_usage", "memory_mb", "request_latency_ms", "error_rate"}
    try:
        anomaly_count = int(output.get("anomaly_count", 0))
        anomalies = output.get("anomalies", [])
        metrics_analyzed = int(output.get("metrics_analyzed", 0))
        if anomaly_count < min(3, expected["anomaly_count"]):
            return False
        if metrics_analyzed <= 0 or len(anomalies) == 0:
            return False
        for item in anomalies:
            if not isinstance(item, dict):
                return False
            if item.get("metric") not in valid_metrics:
                return False
            if "timestamp" not in item:
                return False
        return True
    except (TypeError, ValueError):
        return False


def _validate_t14_recurrence(output: dict, *, expected: dict) -> bool:
    try:
        if int(output.get("department_count", -1)) != expected["department_count"]:
            return False
        spending = output.get("department_spending", {})
        tol = expected["spending_tolerance_pct"]
        for dept, exp in expected["department_spending"].items():
            actual = spending.get(dept)
            if actual is None:
                return False
            if abs(float(actual) - exp) / max(abs(exp), 1.0) > (tol / 100.0):
                return False
        return (
            int(output.get("over_budget_count", -1)) == expected["over_budget_count"]
            and int(output.get("flagged_transactions", -1)) == expected["flagged_transactions"]
        )
    except (TypeError, ValueError, AttributeError):
        return False


def _validate_t15_recurrence(output: dict, *, expected: dict) -> bool:
    try:
        return (
            int(output.get("n_days", -1)) == expected["n_days"]
            and int(output.get("metrics_aggregated", -1)) == expected["metrics_aggregated"]
            and int(output.get("alert_days", -1)) == expected["alert_days"]
        )
    except (TypeError, ValueError):
        return False


def _validate_t20_recurrence(output: dict, *, expected: dict) -> bool:
    try:
        total = int(output.get("total_applications", -1))
        eligible = int(output.get("eligible_count", -1))
        ineligible = int(output.get("ineligible_count", -1))
        rules_applied = int(output.get("rules_applied", -1))
        return (
            total == expected["total_applications"]
            and eligible == expected["eligible_count"]
            and ineligible == expected["ineligible_count"]
            and rules_applied == expected["rules_applied"]
            and eligible + ineligible == total
        )
    except (TypeError, ValueError):
        return False


_VALIDATOR_FACTORY_MAP = {
    "T04": lambda expected, _task: {
        "validator": partial(_validate_t04_recurrence, expected=expected),
    },
    "T11": lambda expected, _task: (
        {
            "validator": partial(_validate_t11_recurrence, expected=expected),
            "file_validator": partial(_validate_t11_recurrence_file, expected=expected),
        }
    ),
    "T13": lambda expected, _task: {
        "validator": partial(_validate_t13_recurrence, expected=expected),
    },
    "T14": lambda expected, _task: {
        "validator": partial(_validate_t14_recurrence, expected=expected),
    },
    "T15": lambda expected, _task: {
        "validator": partial(_validate_t15_recurrence, expected=expected),
    },
    "T20": lambda expected, _task: {
        "validator": partial(_validate_t20_recurrence, expected=expected),
    },
}


def build_recurrence_tasks(
    *,
    task_ids: list[str] | None = None,
    variants: list[str] | None = None,
) -> list[TaskDefinition]:
    """Return recurrence task instances from the generated manifest."""
    rows = _manifest_rows()
    selected_tasks = set(task_ids or [])
    selected_variants = set(variants or [])
    built: list[TaskDefinition] = []
    for row in rows:
        task_family_id = row["task_family_id"]
        variant_id = row["variant_id"]
        if selected_tasks and task_family_id not in selected_tasks:
            continue
        if selected_variants and variant_id not in selected_variants:
            continue
        canonical = get_task(task_family_id)
        fixture_paths = _rebind_fixture_paths(
            canonical,
            [ROOT / rel for rel in row["fixture_files"]],
        )
        validator_updates = _VALIDATOR_FACTORY_MAP[task_family_id](
            row["expected_values"],
            canonical,
        )
        built.append(
            replace(
                canonical,
                task_prompt=_rewrite_prompt(canonical, fixture_paths),
                fixture_paths=fixture_paths,
                task_family_id=task_family_id,
                task_instance_id=row["task_instance_id"],
                variant_id=variant_id,
                benchmark_track=row.get("benchmark_track", "recurrence"),
                **validator_updates,
            )
        )
    return built
